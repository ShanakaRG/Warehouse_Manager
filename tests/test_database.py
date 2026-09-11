"""Business-rule tests. Run with:  pytest"""
from datetime import date

import pytest
from openpyxl import load_workbook

from app.database import IN, OUT, Database, InsufficientStockError, ValidationError
from app.export import export_all


@pytest.fixture
def db():
    d = Database(":memory:")
    d.set_setting("username", "tester")
    d.set_setting("company_name", "ACME")
    yield d
    d.close()


@pytest.fixture
def cement(db):
    return db.add_material("Cement", "bag", min_stock=10)


def receive(db, mid, qty, **kw):
    return db.add_transaction(IN, mid, date.today(), qty, party_name="Supplier A", **kw)


def issue(db, mid, qty, **kw):
    return db.add_transaction(OUT, mid, date.today(), qty, party_name="Site 1", **kw)


def test_incoming_increases_balance(db, cement):
    receive(db, cement, 50)
    assert db.get_balance(cement) == 50


def test_outgoing_cannot_exceed_balance(db, cement):
    receive(db, cement, 20)
    with pytest.raises(InsufficientStockError):
        issue(db, cement, 20.001)
    assert db.get_balance(cement) == 20


def test_issue_exact_balance_leaves_zero(db, cement):
    receive(db, cement, 20)
    issue(db, cement, 20)
    assert db.get_balance(cement) == 0


def test_cannot_issue_from_empty_stock(db, cement):
    with pytest.raises(InsufficientStockError):
        issue(db, cement, 1)


def test_quantity_must_be_positive(db, cement):
    with pytest.raises(ValidationError):
        receive(db, cement, 0)
    with pytest.raises(ValidationError):
        receive(db, cement, -5)


def test_delete_consumed_incoming_is_blocked(db, cement):
    tid = receive(db, cement, 10)
    issue(db, cement, 8)
    with pytest.raises(InsufficientStockError):
        db.delete_transaction(tid)
    assert db.get_balance(cement) == 2


def test_edit_incoming_below_consumed_is_blocked(db, cement):
    tid = receive(db, cement, 10)
    issue(db, cement, 8)
    with pytest.raises(InsufficientStockError):
        db.update_transaction(tid, cement, date.today(), 7, party_name="Supplier A")
    db.update_transaction(tid, cement, date.today(), 8, party_name="Supplier A")
    assert db.get_balance(cement) == 0


def test_edit_outgoing_can_reuse_its_own_quantity(db, cement):
    receive(db, cement, 10)
    oid = issue(db, cement, 10)
    db.update_transaction(oid, cement, date.today(), 10, party_name="Site 2")  # no change in qty
    with pytest.raises(InsufficientStockError):
        db.update_transaction(oid, cement, date.today(), 11, party_name="Site 2")


def test_moving_outgoing_to_another_material(db, cement):
    sand = db.add_material("Sand", "m3")
    receive(db, cement, 5)
    oid = issue(db, cement, 5)
    with pytest.raises(InsufficientStockError):   # sand has nothing to give
        db.update_transaction(oid, sand, date.today(), 5, party_name="Site 1")
    receive(db, sand, 5)
    db.update_transaction(oid, sand, date.today(), 5, party_name="Site 1")
    assert db.get_balance(cement) == 5 and db.get_balance(sand) == 0


def test_float_precision(db, cement):
    receive(db, cement, 0.1)
    receive(db, cement, 0.2)
    issue(db, cement, 0.3)       # 0.1 + 0.2 != 0.3 in raw floats; must still be allowed
    assert db.get_balance(cement) == 0


def test_duplicate_material_names_case_insensitive(db, cement):
    with pytest.raises(ValidationError):
        db.add_material("CEMENT")


def test_material_with_entries_cannot_be_deleted(db, cement):
    receive(db, cement, 1)
    with pytest.raises(ValidationError):
        db.delete_material(cement)


def test_extra_fields_round_trip(db, cement):
    tid = receive(db, cement, 3, extra_fields={"Batch no": "B-77", "  ": "ignored"})
    assert db.get_transaction(tid)["extra_fields"] == {"Batch no": "B-77"}
    assert db.extra_field_names() == ["Batch no"]


def test_stock_status(db, cement):
    receive(db, cement, 10)
    row = next(b for b in db.stock_balance() if b["id"] == cement)
    assert row["status"] == "Low stock"


def test_audit_log_records_changes(db, cement):
    tid = receive(db, cement, 5)
    db.delete_transaction(tid)
    actions = [a["action"] for a in db.list_audit()]
    assert actions[:2] == ["DELETE", "CREATE"]


def test_export_all(db, cement, tmp_path):
    receive(db, cement, 12, extra_fields={"Batch no": "B-1"})
    issue(db, cement, 2)
    path = export_all(db, tmp_path / "out.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Stock balance", "Incoming", "Outgoing", "Change history"]
    assert wb["Stock balance"]["E2"].value == 10
    assert "Batch no" in [c.value for c in wb["Incoming"][1]]
