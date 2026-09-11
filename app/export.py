"""Excel export (no GUI dependency, so it is unit-tested)."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .database import IN, OUT, Database
from .utils import STATUS_LOW, STATUS_OUT
from .version import __app_name__, __version__

HEADER_FILL = PatternFill("solid", fgColor="243447")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="243447")
STATUS_FILLS = {
    STATUS_LOW: PatternFill("solid", fgColor="FFE8A3"),
    STATUS_OUT: PatternFill("solid", fgColor="F8C4BD"),
}
QTY_FORMAT = "#,##0.###"
DATE_FORMAT = "dd/mm/yyyy"


def _to_date(iso: str):
    try:
        return datetime.strptime(iso, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return iso


def _write_table(ws, headers: list[str], rows: list[list], qty_cols=(), date_cols=()):
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    header_row = ws.max_row
    for row in rows:
        ws.append(row)
    for col_idx in qty_cols:
        for (cell,) in ws.iter_rows(min_row=header_row + 1, min_col=col_idx, max_col=col_idx):
            cell.number_format = QTY_FORMAT
    for col_idx in date_cols:
        for (cell,) in ws.iter_rows(min_row=header_row + 1, min_col=col_idx, max_col=col_idx):
            cell.number_format = DATE_FORMAT
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{ws.max_row}"
    _autofit(ws)


def _autofit(ws):
    for column in ws.columns:
        width = 8
        for cell in column:
            value = cell.value
            if isinstance(value, (date, datetime)):
                length = 10
            else:
                length = max((len(line) for line in str(value or "").split("\n")), default=0)
            width = max(width, length + 2)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(width, 50)


def _txn_sheet(wb, title: str, txns: list[dict], include_type: bool):
    ws = wb.create_sheet(title)
    extra_keys = sorted({k for t in txns for k in t["extra_fields"]}, key=str.lower)
    headers = ["Entry ID"] + (["Type"] if include_type else []) + [
        "Date", "Material", "Quantity", "Unit", "Supplier / Issued to", "Party details",
        "Reference no.", "Remarks", *extra_keys, "Entered by", "Entered at", "Last edited by",
    ]
    rows = []
    for t in txns:
        row = [t["id"]]
        if include_type:
            row.append("Incoming" if t["txn_type"] == IN else "Outgoing")
        row += [
            _to_date(t["txn_date"]), t["material_name"], t["quantity"], t["unit"],
            t["party_name"], t["party_details"], t["reference_no"], t["remarks"],
            *[t["extra_fields"].get(k, "") for k in extra_keys],
            t["created_by"], t["created_at"].replace("T", " "), t.get("updated_by") or "",
        ]
        rows.append(row)
    offset = 1 if include_type else 0
    _write_table(ws, headers, rows, qty_cols=(4 + offset,), date_cols=(2 + offset,))


def _balance_sheet(wb, balances: list[dict]):
    ws = wb.create_sheet("Stock balance")
    headers = ["Material", "Unit", "Total in", "Total out", "Balance", "Minimum stock", "Status"]
    rows = [[b["name"], b["unit"], b["total_in"], b["total_out"], b["balance"],
             b["min_stock"], b["status"]] for b in balances]
    _write_table(ws, headers, rows, qty_cols=(3, 4, 5, 6))
    for (cell,) in ws.iter_rows(min_row=2, min_col=7, max_col=7):
        fill = STATUS_FILLS.get(cell.value)
        if fill:
            cell.fill = fill


def _summary_sheet(ws, db: Database, title: str):
    ws.title = "Summary"
    ws["A1"] = db.company_name or "Warehouse"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = title
    info = [
        ("Exported by", db.username),
        ("Exported on", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Software", f"{__app_name__} v{__version__}"),
    ]
    for i, (label, value) in enumerate(info, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 40


def export_all(db: Database, path: str | Path) -> Path:
    """Everything: summary, stock balance, incoming, outgoing and change history."""
    wb = Workbook()
    _summary_sheet(wb.active, db, "Full warehouse export")
    _balance_sheet(wb, db.stock_balance())
    _txn_sheet(wb, "Incoming", db.list_transactions(txn_type=IN), include_type=False)
    _txn_sheet(wb, "Outgoing", db.list_transactions(txn_type=OUT), include_type=False)
    audit = wb.create_sheet("Change history")
    _write_table(audit, ["Time", "User", "Action", "Item", "Item ID", "Details"],
                 [[a["ts"].replace("T", " "), a["username"], a["action"], a["entity"],
                   a["entity_id"], a["details"]] for a in db.list_audit()])
    path = Path(path)
    wb.save(path)
    return path


def export_transactions(db: Database, path: str | Path, txns: list[dict],
                        title: str = "Filtered records") -> Path:
    wb = Workbook()
    _summary_sheet(wb.active, db, title)
    _txn_sheet(wb, "Records", txns, include_type=True)
    path = Path(path)
    wb.save(path)
    return path


def export_balance(db: Database, path: str | Path) -> Path:
    wb = Workbook()
    _summary_sheet(wb.active, db, "Stock balance report")
    _balance_sheet(wb, db.stock_balance())
    path = Path(path)
    wb.save(path)
    return path
